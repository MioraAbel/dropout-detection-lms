import docx
from docx.shared import Pt
import os

desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
file_path = os.path.join(desktop, 'Test_Route_Alertes_PFE.docx')

try:
    doc = docx.Document()
    
    # Titre principal
    title = doc.add_heading('Validation de la Route de Gestion des Alertes', 0)
    title.alignment = 1 # Center
    
    # Introduction
    doc.add_heading('1. Objectif du Test', level=1)
    doc.add_paragraph(
        "Ce test final valide le module de gestion des alertes critiques du système (POST et GET /api/alertes). "
        "Ce module est chargé d'avertir de manière proactive l'équipe pédagogique lorsqu'un étudiant est détecté "
        "en situation de décrochage imminent par le modèle de Machine Learning."
    )
    
    # Explication technique
    doc.add_heading('2. Processus Technique et Tolérance aux Pannes (Fault Tolerance)', level=1)
    doc.add_paragraph(
        "L'architecture de cette route a été conçue pour être robuste. Lorsqu'une requête de création d'alerte (POST) est reçue, "
        "l'API effectue deux actions majeures :"
    )
    p = doc.add_paragraph()
    p.add_run("1. Persistance des données : ").bold = True
    p.add_run("L'alerte est insérée en base de données de manière asynchrone pour garantir une historisation (Table 'alert').\n")
    p.add_run("2. Notification externe (Email) : ").bold = True
    p.add_run("Le système tente d'envoyer un e-mail via le module 'flask_mail'.")
    
    doc.add_paragraph(
        "Point fort architectural : Le système intègre un mécanisme de 'try/except' avancé. Ainsi, même si le serveur SMTP "
        "(serveur d'e-mails) est injoignable, l'API ne crashe pas. Elle sauvegarde l'alerte et renvoie un code HTTP 201 (Created) "
        "avec un message de repli indiquant que l'alerte est sauvegardée mais l'e-mail non envoyé."
    )

    # Les résultats (les photos de l'utilisateur)
    doc.add_heading('3. Résultats Obtenus', level=1)
    doc.add_paragraph(
        "Le test d'injection d'une alerte manuelle via le client HTTP a confirmé le bon fonctionnement du mécanisme de "
        "sauvegarde et la robustesse de l'API face à la simulation d'une panne du serveur mail."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN HOPPSCOTCH (POST) JUSTE ICI <--- ]").bold = True
    
    doc.add_paragraph(
        "La requête de consultation (GET) confirme ensuite que la base de données a bien enregistré et restitué "
        "l'alerte formatée pour l'interface utilisateur."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN DU NAVIGATEUR (GET) JUSTE ICI <--- ]").bold = True
    
    # Sauvegarde
    doc.save(file_path)
    print("Le document a été créé avec succès !")
    
except Exception as e:
    print(f"Erreur : {e}")
