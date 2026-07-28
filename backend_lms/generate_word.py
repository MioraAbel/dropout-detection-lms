import docx
from docx.shared import Pt
import os

# Create a new Document
doc = docx.Document()

# Add a title
title = doc.add_heading('Architecture et Utilité du Backend dans le Système d\'Alerte Précoce', 0)
title.alignment = 1 # Center

# Introduction
doc.add_heading('1. Le Backend : Le "Cerveau" de l\'Application', level=1)
doc.add_paragraph(
    "Dans l'architecture de notre système de prédiction de décrochage scolaire, le Backend (développé en Python avec le framework Flask) joue le rôle central de moteur analytique. "
    "Il agit comme une passerelle intelligente entre la base de données (MySQL) contenant les traces d'apprentissage extraites de Moodle, et notre modèle d'Intelligence Artificielle (Stacking Model basé sur Random Forest, XGBoost et SVM)."
)

# Interaction Frontend-Backend
doc.add_heading('2. Séparation Backend / Frontend', level=1)
doc.add_paragraph(
    "Il est crucial de comprendre que l'API n'a pas vocation à être utilisée directement par l'utilisateur final. "
    "L'utilité principale du Backend est d'exposer des données sécurisées sous format JSON à un Frontend (le Dashboard interactif). "
    "C'est le Dashboard qui se charge de consommer ces données pour générer des interfaces visuelles riches (tableaux de bord, graphiques, alertes de couleur), "
    "rendant l'information exploitable pour l'administration et le corps enseignant de la FST de Settat."
)

# Endpoints
doc.add_heading('3. Les Endpoints (Routes de l\'API)', level=1)
p = doc.add_paragraph()
p.add_run("L'API repose principalement sur deux fonctionnalités clés :\n").bold = True
p.add_run("• GET /api/apprenants : ").bold = True
p.add_run("Cette route récupère la liste complète des étudiants depuis la base de données, la fusionne en temps réel avec leurs statistiques Moodle, et interroge le modèle d'IA pour calculer le niveau de risque de chacun.\n")
p.add_run("• POST /api/predictions/predire : ").bold = True
p.add_run("Cette route permet d'analyser un profil étudiant spécifique à la demande. Si le modèle détecte un risque critique de décrochage, le backend archive automatiquement une alerte dans la table de suivi (historique) et met à jour le statut global de l'étudiant dans la base.")

# EDM Thresholds
doc.add_heading('4. Justification des Seuils d\'Alerte (EDM)', level=1)
doc.add_paragraph(
    "En nous inspirant des systèmes d'alerte précoce (Early Warning Systems) et du framework MTSS (Multi-Tiered System of Supports), "
    "nous avons opté pour une classification des risques en trois niveaux, plutôt qu'une décision binaire (0.5), afin de mieux prioriser les interventions pédagogiques :"
)
doc.add_paragraph("• Risque Élevé (>= 70%) : Signaux de décrochage critiques nécessitant une intervention d'urgence.", style='List Bullet')
doc.add_paragraph("• Risque Modéré (40% - 69%) : Baisse de régime ou absence d'activité suspecte (Zone de surveillance).", style='List Bullet')
doc.add_paragraph("• Risque Faible (< 40%) : Profil d'apprentissage sain.", style='List Bullet')

# Save to desktop
desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
file_path = os.path.join(desktop, 'Resume_Architecture_PFE.docx')
doc.save(file_path)

print(f"Document saved to {file_path}")
