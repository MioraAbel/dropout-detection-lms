import docx
from docx.shared import Pt
import os

desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
file_path = os.path.join(desktop, 'Resume_Architecture_PFE.docx')

try:
    # On ouvre le document existant pour ne pas écraser les photos de l'utilisateur
    doc = docx.Document(file_path)
    
    # On ajoute une nouvelle section
    doc.add_heading('5. Validation et Tests d\'Intégration de l\'API', level=1)
    doc.add_paragraph(
        "Afin de s'assurer de la robustesse du Backend avant son interfaçage avec le Frontend, "
        "une série de tests d'intégration a été réalisée. Le workflow ci-dessous démontre la communication "
        "parfaite entre le client HTTP, l'API Flask, le modèle de Machine Learning et la base de données MySQL."
    )
    
    # Etape 1
    doc.add_heading('Étape 1 : Simulation de la Requête Client (Outil Hoppscotch)', level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Description : ").bold = True
    p1.add_run(
        "Pour contourner l'absence d'interface graphique durant la phase de développement, "
        "nous avons utilisé l'outil Hoppscotch pour simuler un client Web. "
        "La requête HTTP de type POST est envoyée à la route ")
    p1.add_run("/api/predictions/predire").italic = True
    p1.add_run(" en intégrant le profil brut d'un étudiant au format JSON.")
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN DE HOPPSCOTCH JUSTE ICI <--- ]").bold = True
    
    # Etape 2
    doc.add_heading('Étape 2 : Réponse de l\'API et Calcul de l\'IA', level=2)
    p2 = doc.add_paragraph()
    p2.add_run("Description : ").bold = True
    p2.add_run(
        "Dès réception du payload JSON, l'API transmet les caractéristiques (features) de l'étudiant "
        "au modèle d'Intelligence Artificielle. Le modèle traite ces statistiques en temps réel et calcule "
        "une probabilité de décrochage. L'API retourne alors une réponse avec un code HTTP 200 (Succès), affichant "
        "le niveau de risque calculé (ex: Élevé, avec un score de 0.995)."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN DU NAVIGATEUR OU DU RESULTAT JSON JUSTE ICI <--- ]").bold = True
    
    # Etape 3
    doc.add_heading('Étape 3 : Persistance des Données dans MySQL', level=2)
    p3 = doc.add_paragraph()
    p3.add_run("Description : ").bold = True
    p3.add_run(
        "Le processus ne s'arrête pas au simple calcul. Comme le montre la capture ci-dessous, "
        "le Backend déclenche automatiquement une requête SQL pour archiver cette alerte. "
        "La table 'prediction' est mise à jour avec l'ID de l'étudiant, le score de risque calculé, "
        "le statut de l'alerte, ainsi qu'un horodatage (Timestamp) précis pour la traçabilité."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN DE VOTRE TERMINAL OU DE PHPMyAdmin JUSTE ICI <--- ]").bold = True

    # Sauvegarde
    doc.save(file_path)
    print("Le document a été mis à jour avec succès !")
    
except Exception as e:
    print(f"Erreur : {e}")
