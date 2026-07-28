import docx
from docx.shared import Pt
import os

desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
file_path = os.path.join(desktop, 'Test_Route_Apprenants_PFE.docx')

try:
    doc = docx.Document()
    
    # Titre principal
    title = doc.add_heading('Validation de la Route Principale : Prédiction de Masse (Batch Prediction)', 0)
    title.alignment = 1 # Center
    
    # Introduction
    doc.add_heading('1. Objectif du Test', level=1)
    doc.add_paragraph(
        "Ce test a pour but de valider le bon fonctionnement de la route API principale de notre application "
        "(GET /api/apprenants/). Cette route est le pilier du futur Dashboard, car elle doit traiter simultanément "
        "des milliers de profils étudiants en temps réel."
    )
    
    # Explication technique
    doc.add_heading('2. Processus Technique de la Route', level=1)
    p = doc.add_paragraph()
    p.add_run("Lorsqu'une requête GET est envoyée à cette URL, le Backend effectue les actions suivantes :\n").bold = True
    p.add_run("1. Extraction SQL : ").bold = True
    p.add_run("Il interroge la base de données MySQL pour récupérer l'identité et les ID de tous les étudiants enregistrés.\n")
    p.add_run("2. Fusion des Traces d'Apprentissage : ").bold = True
    p.add_run("Il lit dynamiquement le dataset issu de Moodle (comportant les features comme le nombre de connexions, la moyenne, etc.) et l'associe aux étudiants de la base de données.\n")
    p.add_run("3. Inférence du Modèle (Stacking) : ").bold = True
    p.add_run("Il transmet toutes ces données d'un seul coup au modèle de Machine Learning pour réaliser une 'Batch Prediction' (Prédiction par lots). L'IA évalue le profil de chaque étudiant et lui attribue une probabilité de décrochage.\n")
    p.add_run("4. Sérialisation JSON : ").bold = True
    p.add_run("Il formate les résultats en JSON en catégorisant les risques (Élevé, Modéré, Faible) pour faciliter l'affichage futur sur le Frontend.")

    # Les résultats (les photos de l'utilisateur)
    doc.add_heading('3. Résultats Obtenus (Captures d\'écran)', level=1)
    doc.add_paragraph(
        "L'exécution de la requête dans le navigateur confirme que le Backend parvient à traiter et renvoyer les "
        "données de plus de 3000 étudiants de manière fluide et structurée."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN AVEC LES BLOCS (0..99, 100..199, etc.) JUSTE ICI <--- ]").bold = True
    
    doc.add_paragraph(
        "En dépliant les données JSON, on observe que le processus de fusion et de prédiction s'est déroulé avec succès. "
        "Contrairement à une initialisation par défaut, le modèle a bien évalué les véritables statistiques Moodle de chaque étudiant, "
        "ce qui se traduit par des niveaux de risques variés (Élevé, Faible, etc.) parfaitement cohérents avec leurs profils réels."
    )
    
    doc.add_paragraph("[ ---> GLISSEZ LA CAPTURE D'ÉCRAN AVEC LE DETAIL DES ETUDIANTS ET LEURS RISQUES VARIES JUSTE ICI <--- ]").bold = True

    # Sauvegarde
    doc.save(file_path)
    print("Le document a été créé avec succès !")
    
except Exception as e:
    print(f"Erreur : {e}")
